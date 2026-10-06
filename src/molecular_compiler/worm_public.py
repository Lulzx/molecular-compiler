"""M1 adapters for public C. elegans sources used by Phase 0.

Raw files are read from a local cache populated outside this module. Every
source is registered (M1-R5) before conversion, and every approximation that
the canonical schema cannot express directly is listed in the ingestion report.
"""

import json
import re
from hashlib import sha256
from pathlib import Path

import numpy as np

SPECIES = "C. elegans"
CONNECTOME_ID = "cook2019_hermaphrodite"
EXPRESSION_ID = "cengen_taylor2021"
PAIRS_ID = "beets2023_peptide_gpcr"
SIGNS_ID = "fenyves2020_polarity"
ATLAS_ID = "randi2023_signal_propagation"
SEQUENCES_ID = "uniprot_celegans"
PLM_ID = "esm2_t33_650M_UR50D"
REFERENCE_ID = "cect_c302_reference"
WORMBASE_ID = "wormbase_ws286"
PEPTIDE_CONNECTOME_ID = "ripoll_sanchez2023_peptide_connectome"

# Relative to the cache root. Hashes are recorded at ingestion, not assumed.
SOURCE_FILES = {
    CONNECTOME_ID: "cect/cect/data/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",
    REFERENCE_ID: "cect/cect/cache/Cook2019HermReader.json",
    EXPRESSION_ID: "wna/wormneuroatlas/data/cengen.h5",
    PAIRS_ID: "wna/wormneuroatlas/data/deorphanization_media_6.csv",
    SIGNS_ID: "wna/wormneuroatlas/data/journal.pcbi.1007974.s003.xlsx",
    ATLAS_ID: "wna/wormneuroatlas/data/funatlas.h5",
    WORMBASE_ID: "wna/wormneuroatlas/data/c_elegans.PRJNA13758.WS286.geneIDs.txt",
    PEPTIDE_CONNECTOME_ID: "cect/cect/data/01022024_neuropeptide_connectome_long_range_model.csv",
}

# CeNGEN resolves some canonical classes into subclasses and merges VD/DD.
CENGEN_TO_CANONICAL = {
    "ASEL": "ASE",
    "ASER": "ASE",
    "AWC_OFF": "AWC",
    "AWC_ON": "AWC",
    "IL2_DV": "IL2",
    "IL2_LR": "IL2",
    "RMD_DV": "RMD",
    "RMD_LR": "RMD",
    "RME_DV": "RME",
    "RME_LR": "RME",
    "DA9": "DA",
    "DB01": "DB",
    "VA12": "VA",
    "VB01": "VB",
    "VB02": "VB",
    "VC_4_5": "VC",
}
AWC_UNIT = "AWC_mean"


def file_hash(path):
    digest = sha256()
    with Path(path).open("rb") as stream:
        while chunk := stream.read(1 << 20):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_name(name):
    """Cook writes VA01; the atlas and CeNGEN-derived tables write VA1."""
    name = str(name).strip()
    match = re.fullmatch(r"([A-Z]+)0*(\d+)", name)
    return f"{match.group(1)}{match.group(2)}" if match else name


def cengen_unit(neuron, classes):
    """Expression unit for a neuron. AWC left/right identity is stochastic."""
    n = normalize_name(neuron)
    if n in {"AWCL", "AWCR"}:
        return AWC_UNIT
    if n in {"ASEL", "ASER"}:
        return n
    special = {
        r"IL2[DV][LR]": "IL2_DV",
        r"IL2[LR]": "IL2_LR",
        r"RMD[DV][LR]": "RMD_DV",
        r"RMD[LR]": "RMD_LR",
        r"RME[DV]": "RME_DV",
        r"RME[LR]": "RME_LR",
        r"DA9": "DA9",
        r"DA[1-8]": "DA",
        r"DB1": "DB01",
        r"DB[2-7]": "DB",
        r"VA12": "VA12",
        r"VA([1-9]|1[01])": "VA",
        r"VB1": "VB01",
        r"VB2": "VB02",
        r"VB([3-9]|1[01])": "VB",
        r"VC[45]": "VC_4_5",
        r"VC[1236]": "VC",
        r"(VD|DD)\d+": "VD_DD",
    }
    for pattern, unit in special.items():
        if re.fullmatch(pattern, n):
            return unit
    candidates = [n]
    if n[-1:] in "LR":
        candidates.append(n[:-1])
    for c in list(candidates):
        if c[-1:] in "DV" and len(c) > 3:
            candidates.append(c[:-1])
    candidates += [re.sub(r"\d+$", "", c) for c in list(candidates)]
    for c in candidates:
        if c in classes:
            return c
    raise ValueError(f"no CeNGEN class for neuron {neuron}")


def canonical_class(neuron, classes):
    """The 118 classes of the pre-registered partition (Section 4.6)."""
    n = normalize_name(neuron)
    if re.fullmatch(r"VD\d+", n):
        return "VD"
    if re.fullmatch(r"DD\d+", n):
        return "DD"
    unit = cengen_unit(n, classes)
    if unit == AWC_UNIT:
        return "AWC"
    return CENGEN_TO_CANONICAL.get(unit, unit)


_FAMILIES = [
    ("innexin", r"inx-\d+|unc-7|unc-9|eat-5"),
    ("peptide", r"(flp|nlp|ins|pdf|ntc|snet|capa|trh)-\d+"),
    (
        "gpcr",
        (
            r"(npr|frpr|dmsr|tkr|ckr|nmur|sprr|gnrr|pdfr|trhr|ntr|fshr|dop|ser|octr"
            r"|tyra|mgl|gar|gbb)-\d+|aex-2|egl-6|seb-3"
        ),
    ),
    (
        "receptor",
        (
            r"(glr|nmr|glc|acr|acc|lgc|ggr)-\d+|avr-1[45]|lev-1|unc-29|unc-38|unc-63"
            r"|deg-3|des-2|eat-2|gab-1|unc-49|exp-1|mod-1|lev-8"
        ),
    ),
    (
        "channel",
        (
            r"egl-19|unc-2|cca-1|nca-[12]|shk-1|shl-1|(shw|kvs|kqt|kcnl|irk|twk|cng|ocr"
            r"|trp|gtl|del|asic|egas|clh|best|acd|tmc)-\d+|exp-2|egl-36|egl-2|unc-103"
            r"|slo-[12]|sup-9|unc-58|egl-23|tax-[24]|osm-9|trpa-1|gon-2|pkd-2|lov-1"
            r"|deg-1|mec-4|mec-10|unc-8|flr-1|pezo-1|unc-105"
        ),
    ),
    (
        "transporter",
        (
            r"(kcc|abts|vglu|glt|ncx)-\d+|nkcc-1|eat-4|unc-47|unc-17|cat-1|snf-3"
            r"|snf-11|mod-5|dat-1|eat-6"
        ),
    ),
    (
        "effector",
        (
            r"egl-30|goa-1|gsa-1|(gpa|pde|gpb|gpc|rgs)-\d+|egl-8|dgk-1|acy-1|pkc-1"
            r"|kin-2|unc-31|egl-10|eat-16"
        ),
    ),
    ("other", r"unc-25|cha-1|tph-1|cat-2|tdc-1|tbh-1"),
]
_FAMILY_PATTERNS = [(c, re.compile(rf"(?:{p})")) for c, p in _FAMILIES]


def classify_gene(name):
    """Molecule class from a curated gene-family list, or None if not modeled."""
    for molecule_class, pattern in _FAMILY_PATTERNS:
        if pattern.fullmatch(name):
            return molecule_class
    return None


def gene_family(name):
    """Family label used by the Section 12.1 recovery check (e.g. glr, twk)."""
    return re.sub(r"-\d+[a-z]?$", "", name)


def _xlsx_matrix(path, sheet):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[sheet].iter_rows(values_only=True))
    columns = [normalize_name(c) if c else None for c in rows[2][3:]]
    matrix = {}
    for row in rows[3:]:
        source = row[2]
        if not source:
            continue
        for target, value in zip(columns, row[3:]):
            if target and value not in (None, 0, ""):
                matrix[(normalize_name(source), target)] = float(value)
    return matrix


def read_cook2019(path, neurons):
    """Chemical (pre, post) section counts and symmetric gap counts."""
    keep = set(neurons)
    chemical = {
        k: v
        for k, v in _xlsx_matrix(path, "hermaphrodite chemical").items()
        if k[0] in keep and k[1] in keep
    }
    raw_gap = _xlsx_matrix(path, "hermaphrodite gap jn symmetric")
    gap = {}
    for (a, b), value in raw_gap.items():
        if a in keep and b in keep and a != b:
            pair = tuple(sorted((a, b)))
            if pair in gap and gap[pair] != value:
                raise ValueError(f"asymmetric corrected gap value for {pair}")
            gap[pair] = value
    return chemical, gap


def read_c302_reference(path, neurons):
    """Counts from the c302/ConnectomeToolbox reader cache (M1-R6)."""
    payload = json.loads(Path(path).read_text())
    nodes = [normalize_name(n) for n in payload["nodes"]]
    keep = [i for i, n in enumerate(nodes) if n in set(neurons)]
    chem = np.asarray(payload["connections"]["Generic_CS"])[np.ix_(keep, keep)]
    gap = np.asarray(payload["connections"]["Generic_GJ"])[np.ix_(keep, keep)]
    gap_pairs = int(np.sum(np.triu(np.maximum(gap, gap.T) > 0, 1)))
    return {
        "neuron_ids": sorted(nodes[i] for i in keep),
        "chemical_edges": int(np.sum(chem > 0)),
        "gap_contacts": gap_pairs,
        "chemical_weight": float(chem.sum()),
    }


def read_cengen(path, threshold=2):
    import h5py

    with h5py.File(path, "r") as h5:

        def text(key):
            return [x.decode() if isinstance(x, bytes) else str(x) for x in h5[key][:]]

        classes = text("neuron_ids")
        names = text(f"gene_names_th{threshold}")
        wbids = text(f"gene_wbids_th{threshold}")
        tpm = np.asarray(h5[f"tpm_th{threshold}"][:], dtype=np.float64)
    if tpm.shape != (len(classes), len(names)) or np.any(tpm < 0):
        raise ValueError("CeNGEN matrix shape or values invalid")
    return classes, names, wbids, tpm


def read_beets_pairs(path):
    """Precursor-gene x receptor-gene pairs; most potent mature peptide kept."""
    import pandas as pd

    table = pd.read_csv(path).dropna(subset=["GPCR name", "Peptide", "EC50 (M)"])
    pairs, sequences = {}, {}
    for _, row in table.iterrows():
        peptide = str(row["Peptide"]).replace("pyro", "")
        peptide = re.match(r"([A-Za-z]+-\d+)", peptide).group(1).lower()
        receptor = re.match(r"([A-Za-z]+-\d+)", str(row["GPCR name"])).group(1).lower()
        sequences[receptor] = re.sub(r"-\d+$", "", str(row["GPCR ID"]))
        ec50 = float(row["EC50 (M)"]) * 1e9
        if not np.isfinite(ec50) or ec50 <= 0:
            continue
        key = (peptide, receptor)
        pairs[key] = min(ec50, pairs.get(key, np.inf))
    return [
        {
            "peptide": p,
            "receptor": r,
            "receptor_sequence_name": sequences[r],
            "ec50_nM": v,
        }
        for (p, r), v in sorted(pairs.items())
    ]


def read_wormbase_ids(path):
    """Public and sequence names to WBGene identifiers (WormBase WS286)."""
    import csv

    mapping = {}
    with Path(path).open() as stream:
        for row in csv.reader(stream):
            if len(row) >= 4 and row[1].startswith("WBGene"):
                for name in row[2:4]:
                    if name:
                        mapping.setdefault(name, row[1])
    return mapping


def read_fenyves(path):
    """Transmitter identity per neuron and ionotropic receptor polarity classes."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    transmitters = {}
    for row in list(workbook["1. NT expr"].iter_rows(values_only=True))[1:]:
        if row[0]:
            transmitters[normalize_name(row[0])] = [
                str(x) for x in row[1:3] if x not in (None, "", "nan")
            ]
    receptor_rows = list(workbook["2. Receptor gene table"].iter_rows(values_only=True))
    header = receptor_rows[0][:6]
    receptors = {}
    for column, label in enumerate(header):
        transmitter, polarity = label.split()
        genes = [r[column] for r in receptor_rows[1:] if r[column]]
        receptors.setdefault(transmitter, {})[
            "excitatory" if polarity == "Pos" else "inhibitory"
        ] = sorted(str(g) for g in genes)
    return transmitters, receptors


def read_signal_propagation(path):
    """Randi et al. response atlas: [responder, stimulated] matrices per strain."""
    import h5py

    result = {}
    with h5py.File(path, "r") as h5:
        names = [normalize_name(n.decode()) for n in h5["neuron_ids"][:]]
        names = [
            "AWCL" if n == "AWCOF" else ("AWCR" if n == "AWCON" else n) for n in names
        ]
        for strain in ("wt", "unc31"):
            group = h5[strain]
            trials = group["dFF_all"][:]
            result[strain] = {
                "dff": np.asarray(group["dFF"][:], dtype=np.float64),
                "q": np.asarray(group["q"][:], dtype=np.float64),
                "q_eq": np.asarray(group["q_eq"][:], dtype=np.float64),
                "occurrences": np.asarray(group["occ1"][:], dtype=np.int64),
                "trials": [
                    [
                        np.asarray(trials[i, j], dtype=np.float64)
                        for j in range(len(names))
                    ]
                    for i in range(len(names))
                ],
            }
        compiled = h5.attrs["time_compiled"].decode()
    return names, result, compiled


def _uniprot(params):
    import urllib.parse
    import urllib.request

    url = "https://rest.uniprot.org/uniprotkb/stream?" + urllib.parse.urlencode(params)
    with urllib.request.urlopen(url, timeout=180) as response:
        return response.read().decode(), response.headers.get("X-UniProt-Release")


def fetch_uniprot(gene_names, cache_path, aliases=None, batch=40):
    """Canonical and isoform sequences for named C. elegans genes.

    Each gene is queried under its CeNGEN name and any WormBase aliases (current
    public name, sequence name), matched against UniProt gene and ORF names.
    Only genes not yet attempted with the same aliases trigger network requests;
    results are cached.
    """
    cache_path = Path(cache_path)
    aliases = aliases or {}
    payload = (
        json.loads(cache_path.read_text())
        if cache_path.exists()
        else {"release": None, "genes": {}, "attempted": {}}
    )
    attempted = payload.setdefault("attempted", {})
    for gene in payload["genes"]:
        attempted.setdefault(gene, [gene])
    queries = {g: sorted({g, *aliases.get(g, ())}) for g in gene_names}
    pending = [
        g
        for g in gene_names
        if g not in payload["genes"] and attempted.get(g) != queries[g]
    ]
    records = {}
    for start in range(0, len(pending), batch):
        chunk = pending[start : start + batch]
        names = sorted({n for g in chunk for n in queries[g]})
        query = (
            "("
            + " OR ".join(f"gene_exact:{n}" for n in names)
            + ") AND organism_id:6239"
        )
        text, payload["release"] = _uniprot(
            {
                "query": query,
                "format": "tsv",
                "fields": "accession,reviewed,gene_names,gene_orf,sequence",
            }
        )
        for line in text.splitlines()[1:]:
            accession, reviewed, entry_names, orfs, sequence = line.split("\t")
            found = set(entry_names.split()) | set(orfs.split())
            for gene in chunk:
                if found & set(queries[gene]):
                    records.setdefault(gene, {})[accession] = {
                        "reviewed": reviewed == "reviewed",
                        "canonical": sequence,
                        "isoforms": {},
                    }
        for gene in chunk:
            attempted[gene] = queries[gene]
    accessions = sorted({a for entries in records.values() for a in entries})
    isoforms = {}
    for start in range(0, len(accessions), 100):
        chunk = accessions[start : start + 100]
        text, _ = _uniprot(
            {
                "query": " OR ".join(f"accession:{a}" for a in chunk),
                "format": "fasta",
                "includeIsoform": "true",
            }
        )
        for entry in text.split(">")[1:]:
            header, *lines = entry.splitlines()
            accession = header.split("|")[1]
            if "-" in accession:
                isoforms.setdefault(accession.split("-")[0], {})[accession] = "".join(
                    lines
                )
    for gene, entries in records.items():
        for accession, row in entries.items():
            row["isoforms"] = isoforms.get(accession, {})
        payload["genes"][gene] = entries
    if pending or not cache_path.exists():
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(json.dumps(payload))
    return payload


def select_sequences(uniprot):
    """Prefer reviewed entries, then the longest canonical sequence."""
    selected = {}
    for gene, entries in uniprot["genes"].items():
        candidates = [e for e in entries.values() if e["canonical"]]
        if not candidates:
            continue
        best = max(candidates, key=lambda e: (e["reviewed"], len(e["canonical"])))
        isoforms = [s for s in best["isoforms"].values() if s != best["canonical"]]
        selected[gene] = {
            "canonical": best["canonical"],
            "isoforms": sorted(set(isoforms)),
            "reviewed": best["reviewed"],
        }
    return selected


TRANSMITTERS = ("ACh", "Glu", "GABA")
DECLARED_APPROXIMATIONS = {
    "synapse_size": "Cook et al. EM serial-section counts (size units: count)",
    "synapse_geometry": "Per-synapse xyz, path distance and diameter are not in the "
    "source; xyz=0, path_dist_post=0 um and local_diameter_post=0.2 um are declared "
    "defaults (no dendritic attenuation)",
    "soma_xyz": "Somatic positions in um are not available; zeros. Unused for C. "
    "elegans, which has no peptidergic spatial kernel (Section 12.5)",
    "gap_area": "Gap-junction serial-section counts stand in for contact area; "
    "1 count is stored as 1 um2. Areas are relative, not measured",
    "unobserved_contacts": "Only observed gap junctions are listed; membrane contacts "
    "without gap junctions are absent from the source",
    "expression": "CeNGEN threshold-2 (medium) TPM transformed as log1p(TPM); zeros "
    "are below the CeNGEN detection threshold",
    "awc": "AWC ON/OFF identity is stochastic between left and right, so both AWC "
    "neurons use the mean of the AWC_ON and AWC_OFF profiles (unit AWC_mean)",
    "nt_pred": "No EM transmitter predictions; transmitter identity enters through "
    "released_ligands (Fenyves et al. 2020 table)",
    "ortholog_group": "Not mapped (M1-R2: retained as null)",
    "gene_selection": "Curated signaling families (worm_public._FAMILIES), Beets "
    "pair genes and Fenyves receptor genes, restricted to genes in CeNGEN",
}


def register_sources(register, cache_root):
    """M1-R5: every source must be registered; hashes are measured here."""
    root = Path(cache_root)
    records, hashes = {}, {}
    for dataset_id, relative in SOURCE_FILES.items():
        records[dataset_id] = register.require(dataset_id)
        hashes[dataset_id] = file_hash(root / relative)
    return records, hashes


def build_worm_project(
    cache_root, output, register, embedder=None, checkpoint_hash=None, device="cpu"
):
    """Convert registered public worm sources into canonical Phase 0 inputs."""
    from dataclasses import asdict

    from . import data
    from .embeddings import embed_genes, esm2_embedder
    from .provenance import content_hash, inherit

    root, output = Path(cache_root), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    records, hashes = register_sources(register, root)
    for extra in (SEQUENCES_ID, PLM_ID):
        records[extra] = register.require(extra)
    transmitters, receptor_polarity = read_fenyves(root / SOURCE_FILES[SIGNS_ID])
    neurons = sorted(transmitters)
    chemical, gap = read_cook2019(root / SOURCE_FILES[CONNECTOME_ID], neurons)
    reference = read_c302_reference(root / SOURCE_FILES[REFERENCE_ID], neurons)
    classes, names, wbids, tpm = read_cengen(root / SOURCE_FILES[EXPRESSION_ID])
    wormbase = read_wormbase_ids(root / SOURCE_FILES[WORMBASE_ID])
    pairs = read_beets_pairs(root / SOURCE_FILES[PAIRS_ID])

    units = list(classes)
    profiles = np.log1p(tpm)
    awc = [units.index("AWC_OFF"), units.index("AWC_ON")]
    profiles = np.vstack([np.delete(profiles, awc, axis=0), profiles[awc].mean(axis=0)])
    units = [u for i, u in enumerate(units) if i not in awc] + [AWC_UNIT]

    by_wbid = {wb: i for i, wb in enumerate(wbids)}
    forced = {}
    dropped_pairs = []
    for pair in pairs:
        peptide = wormbase.get(pair["peptide"])
        receptor = wormbase.get(pair["receptor"]) or wormbase.get(
            pair["receptor_sequence_name"]
        )
        if peptide in by_wbid and receptor in by_wbid:
            forced[peptide], forced[receptor] = "peptide", "gpcr"
            pair["peptide_wbid"], pair["receptor_wbid"] = peptide, receptor
        else:
            dropped_pairs.append(pair)
    receptor_genes = {
        g
        for polarity in receptor_polarity.values()
        for v in polarity.values()
        for g in v
    }
    for gene in receptor_genes:
        if wormbase.get(gene) in by_wbid:
            forced[wormbase[gene]] = "receptor"
    selected = []
    for i, (name, wb) in enumerate(zip(names, wbids)):
        molecule_class = forced.get(wb) or classify_gene(name)
        if molecule_class and np.any(tpm[:, i] > 0):
            selected.append((i, name, wb, molecule_class))

    names_by_wbid = {}
    for name, wb in wormbase.items():
        names_by_wbid.setdefault(wb, set()).add(name)
    uniprot = fetch_uniprot(
        [s[1] for s in selected],
        root / "uniprot_sequences.json",
        aliases={s[1]: sorted(names_by_wbid.get(s[2], ())) for s in selected},
    )
    sequences = select_sequences(uniprot)
    missing_sequence = [s[1] for s in selected if s[1] not in sequences]
    selected = [s for s in selected if s[1] in sequences]
    gene_rows = [
        {
            "gene_id": wb,
            "species": SPECIES,
            "protein_seq": sequences[name]["canonical"],
            "isoforms": sequences[name]["isoforms"],
            "molecule_class": molecule_class,
            "name": name,
        }
        for _, name, wb, molecule_class in selected
    ]
    embedding_cache = root / "esm2_embeddings.npz"
    cached = dict(np.load(embedding_cache)) if embedding_cache.exists() else {}
    needed = [
        g
        for g in gene_rows
        if sha256(g["protein_seq"].encode()).hexdigest() not in cached
        or any(sha256(s.encode()).hexdigest() not in cached for s in g["isoforms"])
    ]
    plm_meta = None
    if needed:
        if checkpoint_hash is None:
            raise ValueError("embedding new sequences requires the checkpoint hash")
        embedder = embedder or esm2_embedder(
            root / "esm" / f"{PLM_ID}.pt", device=device
        )
        computed, plm_meta = embed_genes(needed, embedder, checkpoint_hash)
        for gene in needed:
            result = computed[gene["gene_id"]]
            cached[sha256(gene["protein_seq"].encode()).hexdigest()] = result[
                "canonical"
            ]
            for sequence, vector in zip(gene["isoforms"], result["isoforms"]):
                cached[sha256(sequence.encode()).hexdigest()] = vector
        np.savez(embedding_cache, **cached)

    def vector(sequence):
        return cached[sha256(sequence.encode()).hexdigest()]

    genes = data.table(
        "genes",
        [
            {
                "gene_id": g["gene_id"],
                "species": SPECIES,
                "protein_seq": g["protein_seq"],
                "isoforms": g["isoforms"],
                "ortholog_group": None,
                "molecule_class": g["molecule_class"],
                "plm_embedding": vector(g["protein_seq"]).astype(np.float16).tolist(),
            }
            for g in gene_rows
        ],
    )
    gene_index = [s[0] for s in selected]
    expression = data.table(
        "expression",
        [
            {
                "unit_id": unit,
                "unit_level": "type",
                "gene_id": wb,
                "value": float(profiles[u, column]),
                "assay": "scRNA",
                "subcellular": "whole",
            }
            for u, unit in enumerate(units)
            for column, wb in zip(gene_index, (g["gene_id"] for g in gene_rows))
            if profiles[u, column] > 0
        ],
    )
    kept = {g["gene_id"] for g in gene_rows}
    pair_table = data.table(
        "peptide_receptor_pairs",
        [
            {
                "peptide_gene_id": p["peptide_wbid"],
                "receptor_gene_id": p["receptor_wbid"],
                "ec50_nM": float(p["ec50_nM"]),
                "evidence": "in_vitro_screen",
                "source": "Beets et al. 2023, Cell Reports, media 6",
            }
            for p in pairs
            if p.get("peptide_wbid") in kept and p.get("receptor_wbid") in kept
        ],
    )
    ids = {n: i for i, n in enumerate(neurons)}
    neuron_table = data.table(
        "neurons",
        [
            {
                "neuron_id": ids[n],
                "animal_id": records[CONNECTOME_ID].animal_id,
                "type_label": cengen_unit(n, classes),
                "soma_xyz": [0.0, 0.0, 0.0],
                "morphology_ref": None,
                "segmentation_confidence": 1.0,
            }
            for n in neurons
        ],
    )
    synapses = data.table(
        "synapses",
        [
            {
                "synapse_id": s,
                "pre_id": ids[a],
                "post_id": ids[b],
                "size": value,
                "vesicle_count": None,
                "xyz": [0.0, 0.0, 0.0],
                "path_dist_post": 0.0,
                "local_diameter_post": 0.2,
                "compartment_post": 0,
                "nt_pred": None,
                "detection_confidence": 1.0,
            }
            for s, ((a, b), value) in enumerate(sorted(chemical.items()))
        ],
    )
    contacts = data.table(
        "contacts",
        [
            {
                "i_id": ids[a],
                "j_id": ids[b],
                "area": value,
                "gap_junction_observed": True,
            }
            for (a, b), value in sorted(gap.items())
        ],
    )
    gene_by_name = {g["name"]: g["gene_id"] for g in gene_rows}
    receptor_ligands = {
        gene_by_name[g]: transmitter
        for transmitter, polarity in receptor_polarity.items()
        for v in polarity.values()
        for g in v
        if g in gene_by_name
    }
    released = {str(ids[n]): transmitters[n] for n in neurons if transmitters[n]}
    explanation = (
        "Neuron set restricted to the 302 hermaphrodite neurons of Fenyves et al. "
        "2020; muscles, glia and end organs in Cook et al. SI 5 are excluded. "
        "The independent parse of the corrected SI 5 workbook and the c302 "
        "ConnectomeToolbox reader cache agree on neurons and edge counts."
    )
    c302 = {k: reference[k] for k in ("neuron_ids", "chemical_edges", "gap_contacts")}
    c302["neuron_ids"] = [ids[n] for n in c302["neuron_ids"]]
    connectome = data.ConnectomeDataset(
        CONNECTOME_ID,
        neuron_table,
        synapses,
        contacts,
        register,
        c302_reference=c302,
        c302_explanation=explanation,
    )
    molecular = data.MolecularDataset(
        EXPRESSION_ID,
        genes,
        expression,
        pair_table,
        register,
        released_ligands=released,
        receptor_ligands=receptor_ligands,
    )
    graph = data.prepare(connectome, molecular)
    inputs = [records[k] for k in sorted(records)]
    metadata = inherit(inputs, content_hash(graph.metadata["processing_hash"], hashes))
    for name in ("neurons", "synapses", "contacts"):
        data.write_table(
            name, getattr(connectome, name), output / f"{name}.parquet", metadata
        )
    for name in ("genes", "expression", "peptide_receptor_pairs"):
        data.write_table(
            name, getattr(molecular, name), output / f"{name}.parquet", metadata
        )
    names_atlas, atlas, compiled = read_signal_propagation(
        root / SOURCE_FILES[ATLAS_ID]
    )
    write_atlas(output / "signal_propagation.npz", names_atlas, atlas, ids)
    report = {
        "species": SPECIES,
        "source_sha256": hashes,
        "register": [asdict(r) for r in inputs],
        "provenance": metadata,
        "neurons": len(neurons),
        "chemical_edges": len(chemical),
        "chemical_section_total": float(sum(chemical.values())),
        "gap_pairs": len(gap),
        "c302_crosscheck": graph.metadata["c302_crosscheck"],
        "c302_reference_counts": reference
        | {"neuron_ids": len(reference["neuron_ids"])},
        "cengen_units": len(units),
        "canonical_classes": len({canonical_class(n, classes) for n in neurons}),
        "genes": dict(
            sorted(
                {
                    c: sum(g["molecule_class"] == c for g in gene_rows)
                    for c in {g["molecule_class"] for g in gene_rows}
                }.items()
            )
        ),
        "genes_without_sequence": missing_sequence,
        "uniprot_release": uniprot.get("release"),
        "plm": plm_meta
        or {"model": PLM_ID, "checkpoint_hash": checkpoint_hash, "cached": True},
        "isoform_embeddings": sum(len(g["isoforms"]) for g in gene_rows),
        "peptide_receptor_pairs": pair_table.num_rows,
        "pairs_dropped_unmapped": [
            f"{p['peptide']}->{p['receptor']}" for p in dropped_pairs
        ],
        "signal_propagation_compiled": compiled,
        "declared_approximations": DECLARED_APPROXIMATIONS,
    }
    manifest = {
        "connectome_dataset_id": CONNECTOME_ID,
        "molecular_dataset_id": EXPRESSION_ID,
        "released_ligands": released,
        "receptor_ligands": receptor_ligands,
        "c302_reference": c302,
        "c302_explanation": explanation,
        "neuron_names": neurons,
        "canonical_classes": {n: canonical_class(n, classes) for n in neurons},
        "gene_names": {g["gene_id"]: g["name"] for g in gene_rows},
        "receptor_polarity": receptor_polarity,
        "gauge_proxies": {
            # GCaMP6s and the GUR-3/PRDX-2 QF driver both use the rab-3 promoter.
            "rab-3": {
                n: float(
                    profiles[units.index(cengen_unit(n, classes)), names.index("rab-3")]
                )
                for n in neurons
            }
        },
    }
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2))
    (output / "ingestion-report.json").write_text(json.dumps(report, indent=2))
    (output / "data-register.json").write_text(
        json.dumps([asdict(r) for r in inputs], indent=2)
    )
    return graph, report


def write_atlas(path, names, atlas, ids):
    """Response matrices re-indexed to canonical neuron IDs; trials flattened."""
    order = np.array([ids[n] for n in names], dtype=np.int64)
    arrays = {"neuron_ids": order}
    for strain, values in atlas.items():
        for key in ("dff", "q", "q_eq", "occurrences"):
            arrays[f"{strain}/{key}"] = values[key]
        lengths = np.array([[len(t) for t in row] for row in values["trials"]])
        arrays[f"{strain}/trial_counts"] = lengths
        arrays[f"{strain}/trials"] = np.concatenate(
            [t for row in values["trials"] for t in row]
        )
    np.savez_compressed(path, **arrays)


def load_atlas(path):
    """[responder, stimulated] matrices indexed by canonical neuron ID order."""
    raw = np.load(path)
    order = raw["neuron_ids"]
    result = {"neuron_ids": order}
    for strain in ("wt", "unc31"):
        counts = raw[f"{strain}/trial_counts"]
        flat = raw[f"{strain}/trials"]
        offsets = np.concatenate([[0], np.cumsum(counts.ravel())])
        trials = [
            [
                flat[offsets[i * len(order) + j] : offsets[i * len(order) + j + 1]]
                for j in range(len(order))
            ]
            for i in range(len(order))
        ]
        result[strain] = {
            key: raw[f"{strain}/{key}"] for key in ("dff", "q", "q_eq", "occurrences")
        } | {"trials": trials}
    return result


def load_worm_project(directory):
    """Reload canonical Phase 0 inputs and rerun M1-M3 preparation."""
    from . import data
    from .provenance import DataRegister

    root = Path(directory)
    manifest = json.loads((root / "manifest.json").read_text())
    register = DataRegister.load(root / "data-register.json")
    connectome = data.ConnectomeDataset(
        manifest["connectome_dataset_id"],
        *[
            data.read_table(name, root / f"{name}.parquet")
            for name in ("neurons", "synapses", "contacts")
        ],
        register,
        c302_reference=manifest["c302_reference"],
        c302_explanation=manifest["c302_explanation"],
    )
    molecular = data.MolecularDataset(
        manifest["molecular_dataset_id"],
        *[
            data.read_table(name, root / f"{name}.parquet")
            for name in ("genes", "expression", "peptide_receptor_pairs")
        ],
        register,
        released_ligands=manifest["released_ligands"],
        receptor_ligands=manifest["receptor_ligands"],
    )
    return data.prepare(connectome, molecular), manifest
